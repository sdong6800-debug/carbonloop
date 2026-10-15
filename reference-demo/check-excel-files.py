from pathlib import Path
from zipfile import ZipFile
from math import isclose
from openpyxl import load_workbook

folder=Path(__file__).parent/'.test-exports'
for file in folder.glob('*.xlsx'):
    with ZipFile(file) as z:
        assert z.testzip() is None
    book=load_workbook(file,data_only=False)
    assert all(c.data_type!='f' for s in book for row in s for c in row), 'Reports must be snapshots, not executable input text'

book=load_workbook(folder/'comparison.xlsx',data_only=True)
assert book.sheetnames==['结果对比','构件明细','材料因子与来源']
summary={row[0]:row[1:] for row in book.worksheets[0].values}
assert isclose(summary['修改后减修改前'][1],-111.8352,abs_tol=1e-7)
detail=list(book['构件明细'].values)
assert len(detail)==13
assert isclose(sum(r[12] for r in detail[1:] if r[0]=='修改前'),summary['已知部分小计'][0])
assert isclose(sum(r[12] for r in detail[1:] if r[0]=='修改后'),summary['已知部分小计'][1])
assert next(r for r in detail if r[0]=='修改后' and r[1]=='S01')[7]==.18
assert book['构件明细']['F2'].data_type=='n'
assert book['构件明细'].auto_filter.ref
assert book['材料因子与来源']['A2'].value=='01.002'
assert book['材料因子与来源']['A2'].data_type=='s'
assert book['材料因子与来源']['D2'].number_format=='0.000000'

partial=load_workbook(folder/'partial.xlsx',data_only=True)
ps={r[0]:r[1:] for r in partial.worksheets[0].values}
assert ps['修改后减修改前'][1]=='不可比较'
assert any(isinstance(r[12],str) and r[13]!='已计算' for r in list(partial['构件明细'].values)[1:])
cases=load_workbook(folder/'cases.xlsx',data_only=True)
assert all(r[1]=='PORR_18' for r in list(cases['案例明细'].values)[1:])
timeline=load_workbook(folder/'timber.xlsx',data_only=True)
assert list(timeline['时间线结果'].values)[-1][0]==200
literal=load_workbook(folder/'literal.xlsx')
assert literal.active['A2'].value=='=1+1' and literal.active['A2'].data_type=='s'
assert literal.active['B2'].value==0 and literal.active['B3'].value=='未提供'
print('PASS independent XLSX parsing, numeric types, exact totals and delta, missing vs zero, filters, 200-year endpoint, literal formula text.')
